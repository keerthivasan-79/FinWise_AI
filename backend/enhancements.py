import csv
import hashlib
import io
import json
import os
import re
import uuid
from datetime import date, datetime, timedelta

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from werkzeug.utils import secure_filename

from auth import token_required
from db import get_conn, query

enhancements = Blueprint("enhancements", __name__)
RECEIPT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "uploads", "receipts"))
ALLOWED_RECEIPTS = {".jpg", ".jpeg", ".png", ".pdf"}


def money(value):
    try:
        result = float(value)
        return result if result > 0 else 0
    except (TypeError, ValueError):
        return 0


def create_transaction(cur, uid, data, recurring_id=None):
    amount=money(data.get("amount")); aid=data.get("account_id"); typ=data.get("type")
    target=data.get("transfer_to_account") or None
    if not amount or typ not in ("Income", "Expense", "Transfer"):
        raise ValueError("A valid type and positive amount are required")
    cur.execute("SELECT * FROM accounts WHERE id=%s AND user_id=%s FOR UPDATE", (aid,uid))
    account=cur.fetchone()
    if not account: raise ValueError("Account not found")
    if typ in ("Expense","Transfer") and float(account["current_balance"]) < amount:
        raise ValueError("Insufficient balance")
    if typ=="Income":
        cur.execute("UPDATE accounts SET current_balance=current_balance+%s WHERE id=%s",(amount,aid))
    elif typ=="Expense":
        cur.execute("UPDATE accounts SET current_balance=current_balance-%s WHERE id=%s",(amount,aid))
    else:
        if not target or int(target)==int(aid): raise ValueError("Choose different transfer accounts")
        cur.execute("SELECT id FROM accounts WHERE id=%s AND user_id=%s FOR UPDATE",(target,uid))
        if not cur.fetchone(): raise ValueError("Target account not found")
        cur.execute("UPDATE accounts SET current_balance=current_balance-%s WHERE id=%s",(amount,aid))
        cur.execute("UPDATE accounts SET current_balance=current_balance+%s WHERE id=%s",(amount,target))
    cur.execute("""INSERT INTO transactions
        (user_id,account_id,type,category,amount,description,transaction_date,payment_method,
         transfer_to_account,tags,merchant,recurring_transaction_id)
        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
        (uid,aid,typ,data.get("category"),amount,data.get("description"),data.get("transaction_date") or date.today(),
         data.get("payment_method"),target,data.get("tags"),data.get("merchant"),recurring_id))
    return cur.lastrowid


def next_occurrence(current, frequency):
    if frequency=="Weekly": return current+timedelta(days=7)
    if frequency=="Quarterly": return current+timedelta(days=91)
    if frequency=="Yearly":
        try: return current.replace(year=current.year+1)
        except ValueError: return current.replace(year=current.year+1,day=28)
    month=current.month+1; year=current.year+(month>12); month=1 if month>12 else month
    last=(date(year+(month==12),1 if month==12 else month+1,1)-timedelta(days=1)).day
    return date(year,month,min(current.day,last))


@enhancements.get("/api/finance/overview")
@token_required
def finance_overview(user):
    uid=user["id"]
    months=query("""SELECT DATE_FORMAT(transaction_date,'%Y-%m') month,
        SUM(CASE WHEN type='Income' THEN amount ELSE 0 END) income,
        SUM(CASE WHEN type='Expense' THEN amount ELSE 0 END) expense
        FROM transactions WHERE user_id=%s AND transaction_date>=DATE_SUB(CURDATE(),INTERVAL 6 MONTH)
        GROUP BY month ORDER BY month""",(uid,),fetch=True)
    expense_values=[float(x["expense"] or 0) for x in months[-3:]]
    forecast=round(sum(expense_values)/len(expense_values),2) if expense_values else 0
    if len(expense_values)>1:
        forecast=max(0,round(forecast+(expense_values[-1]-expense_values[0])/len(expense_values),2))
    subscriptions=query("""SELECT LOWER(TRIM(COALESCE(merchant,description))) merchant,category,
        ROUND(AVG(amount),2) average_amount,COUNT(*) occurrences,MAX(transaction_date) last_seen
        FROM transactions WHERE user_id=%s AND type='Expense'
        AND transaction_date>=DATE_SUB(CURDATE(),INTERVAL 12 MONTH)
        AND COALESCE(merchant,description) IS NOT NULL
        GROUP BY LOWER(TRIM(COALESCE(merchant,description))),category
        HAVING COUNT(*)>=2 ORDER BY average_amount DESC LIMIT 12""",(uid,),fetch=True)
    unusual=query("""SELECT t.id,t.description,t.category,t.amount,t.transaction_date,a.account_name
        FROM transactions t JOIN accounts a ON a.id=t.account_id
        WHERE t.user_id=%s AND t.type='Expense'
        AND t.amount>(SELECT COALESCE(AVG(amount)+2*STDDEV_POP(amount),999999999)
        FROM transactions WHERE user_id=%s AND type='Expense')
        ORDER BY t.transaction_date DESC LIMIT 8""",(uid,uid),fetch=True)
    totals=query("""SELECT COALESCE(SUM(CASE WHEN type='Income' THEN amount ELSE 0 END),0) income,
        COALESCE(SUM(CASE WHEN type='Expense' THEN amount ELSE 0 END),0) expense
        FROM transactions WHERE user_id=%s AND DATE_FORMAT(transaction_date,'%Y-%m')=DATE_FORMAT(CURDATE(),'%Y-%m')""",
        (uid,),fetch=True,one=True)
    income=float(totals["income"]); expense=float(totals["expense"]); savings=income-expense
    rate=(savings/income*100) if income else 0
    score=max(0,min(100,round(50+rate*.45-(15 if expense>income else 0))))
    query("""INSERT INTO health_score_history(user_id,score,recorded_on) VALUES(%s,%s,CURDATE())
             ON DUPLICATE KEY UPDATE score=VALUES(score)""",(uid,score))
    history=query("SELECT score,recorded_on FROM health_score_history WHERE user_id=%s ORDER BY recorded_on DESC LIMIT 90",(uid,),fetch=True)
    accounts=query("SELECT account_type,current_balance,institution_value FROM accounts WHERE user_id=%s",(uid,),fetch=True)
    goal_savings=query("""SELECT IFNULL(SUM(saved_amount),0) total FROM savings_goals
        WHERE user_id=%s OR household_id IN
        (SELECT household_id FROM household_members WHERE user_id=%s AND status='Active')""",(uid,uid),fetch=True,one=True)["total"]
    assets=sum(float(a["institution_value"] or a["current_balance"] or 0) for a in accounts if a["account_type"] not in ("Loan","Credit Card"))
    assets+=float(goal_savings or 0)
    liabilities=sum(abs(float(a["institution_value"] or a["current_balance"] or 0)) for a in accounts if a["account_type"] in ("Loan","Credit Card"))
    recommendations=[]
    if forecast and income and forecast>income*.8: recommendations.append("Forecast spending is above 80% of income; reduce flexible categories now.")
    if subscriptions: recommendations.append(f"Review {len(subscriptions)} possible recurring subscriptions for unused services.")
    if savings>0: recommendations.append(f"Move up to ₹{savings*.5:,.0f} toward goals while retaining a cash buffer.")
    if not recommendations: recommendations.append("Add more transactions to unlock stronger recommendations.")
    if unusual:
        query("""INSERT INTO notifications(user_id,title,message,level)
            SELECT %s,'Unusual spending detected',%s,'warning' FROM DUAL
            WHERE NOT EXISTS (SELECT 1 FROM notifications WHERE user_id=%s AND title='Unusual spending detected'
            AND DATE(created_at)=CURDATE())""",(uid,f"{len(unusual)} unusually large transaction(s) need review.",uid))
    if forecast and income and forecast>income*.8:
        query("""INSERT INTO notifications(user_id,title,message,level)
            SELECT %s,'Cash-flow forecast warning',%s,'warning' FROM DUAL
            WHERE NOT EXISTS (SELECT 1 FROM notifications WHERE user_id=%s AND title='Cash-flow forecast warning'
            AND DATE(created_at)=CURDATE())""",(uid,f"Forecast expenses are {forecast/income*100:.0f}% of monthly income.",uid))
    return jsonify({"forecast":forecast,"months":months,"subscriptions":subscriptions,"unusual":unusual,
                    "score":score,"scoreHistory":history,"netWorth":assets-liabilities,
                    "assets":assets,"liabilities":liabilities,"recommendations":recommendations})


@enhancements.get("/api/recurring")
@token_required
def recurring_list(user):
    return jsonify(query("""SELECT r.*,a.account_name FROM recurring_transactions r
        JOIN accounts a ON a.id=r.account_id WHERE r.user_id=%s ORDER BY r.next_run""",(user["id"],),fetch=True))


@enhancements.post("/api/recurring")
@token_required
def recurring_add(user):
    d=request.get_json(silent=True) or {}; amount=money(d.get("amount"))
    if not d.get("account_id") or not amount or d.get("frequency") not in ("Weekly","Monthly","Quarterly","Yearly"):
        return jsonify({"error":"Account, amount and frequency are required"}),400
    query("""INSERT INTO recurring_transactions
        (user_id,account_id,type,category,amount,description,payment_method,transfer_to_account,frequency,next_run)
        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
        (user["id"],d["account_id"],d.get("type","Expense"),d.get("category"),amount,d.get("description"),
         d.get("payment_method"),d.get("transfer_to_account") or None,d["frequency"],d.get("next_run") or date.today()))
    return jsonify({"message":"Recurring transaction created"}),201


@enhancements.delete("/api/recurring/<int:item_id>")
@token_required
def recurring_delete(user,item_id):
    query("DELETE FROM recurring_transactions WHERE id=%s AND user_id=%s",(item_id,user["id"]))
    return jsonify({"message":"Recurring transaction removed"})


@enhancements.post("/api/recurring/process")
@token_required
def recurring_process(user):
    result=process_due_recurring(user["id"])
    return jsonify(result)


def process_due_recurring(uid):
    rows=query("SELECT * FROM recurring_transactions WHERE user_id=%s AND active=1 AND next_run<=CURDATE()",(uid,),fetch=True)
    conn=get_conn(); cur=conn.cursor(dictionary=True); created=0; errors=[]
    try:
        for row in rows:
            try:
                tx={**row,"transaction_date":row["next_run"]}
                create_transaction(cur,uid,tx,row["id"])
                nxt=next_occurrence(row["next_run"],row["frequency"])
                cur.execute("UPDATE recurring_transactions SET last_run=%s,next_run=%s WHERE id=%s",(row["next_run"],nxt,row["id"]))
                created+=1
            except ValueError as exc:
                errors.append({"id":row["id"],"error":str(exc)})
        conn.commit()
    except Exception:
        conn.rollback(); raise
    finally:
        cur.close(); conn.close()
    return {"created":created,"errors":errors}


@enhancements.get("/api/transactions/duplicates")
@token_required
def duplicates(user):
    return jsonify(query("""SELECT MIN(id) keep_id,GROUP_CONCAT(id ORDER BY id) ids,transaction_date,
        account_id,type,amount,LOWER(TRIM(COALESCE(description,''))) description,COUNT(*) duplicate_count
        FROM transactions WHERE user_id=%s GROUP BY transaction_date,account_id,type,amount,
        LOWER(TRIM(COALESCE(description,''))) HAVING COUNT(*)>1 ORDER BY transaction_date DESC""",(user["id"],),fetch=True))


@enhancements.patch("/api/transactions/bulk")
@token_required
def bulk_transactions(user):
    d=request.get_json(silent=True) or {}; ids=[int(x) for x in d.get("ids",[])][:200]
    allowed={"category","tags","review_status"}; changes={k:v for k,v in d.get("changes",{}).items() if k in allowed}
    if not ids or not changes: return jsonify({"error":"Select transactions and valid changes"}),400
    sets=",".join(f"{k}=%s" for k in changes); placeholders=",".join(["%s"]*len(ids))
    query(f"UPDATE transactions SET {sets} WHERE user_id=%s AND id IN ({placeholders})",
          tuple(changes.values())+(user["id"],)+tuple(ids))
    return jsonify({"message":"Transactions updated","count":len(ids)})


@enhancements.post("/api/transactions/<int:transaction_id>/splits")
@token_required
def save_splits(user,transaction_id):
    splits=(request.get_json(silent=True) or {}).get("splits",[])
    tx=query("SELECT amount FROM transactions WHERE id=%s AND user_id=%s",(transaction_id,user["id"]),fetch=True,one=True)
    if not tx: return jsonify({"error":"Transaction not found"}),404
    if not splits or abs(sum(money(x.get("amount")) for x in splits)-float(tx["amount"]))>.009:
        return jsonify({"error":"Split amounts must equal the transaction total"}),400
    conn=get_conn(); cur=conn.cursor()
    try:
        cur.execute("DELETE FROM transaction_splits WHERE transaction_id=%s AND user_id=%s",(transaction_id,user["id"]))
        cur.executemany("INSERT INTO transaction_splits(transaction_id,user_id,category,amount) VALUES(%s,%s,%s,%s)",
                        [(transaction_id,user["id"],x.get("category","Other"),money(x.get("amount"))) for x in splits])
        conn.commit()
    finally: cur.close(); conn.close()
    return jsonify({"message":"Transaction split saved"})


@enhancements.post("/api/transactions/<int:transaction_id>/receipt")
@token_required
def upload_receipt(user,transaction_id):
    tx=query("SELECT id FROM transactions WHERE id=%s AND user_id=%s",(transaction_id,user["id"]),fetch=True,one=True)
    upload=request.files.get("file")
    if not tx or not upload: return jsonify({"error":"Transaction and receipt are required"}),400
    ext=os.path.splitext(upload.filename or "")[1].lower()
    if ext not in ALLOWED_RECEIPTS: return jsonify({"error":"Use JPG, PNG or PDF receipts"}),400
    os.makedirs(RECEIPT_DIR,exist_ok=True)
    filename=f"{user['id']}_{transaction_id}_{uuid.uuid4().hex[:10]}{ext}"
    upload.save(os.path.join(RECEIPT_DIR,secure_filename(filename)))
    query("UPDATE transactions SET receipt_path=%s WHERE id=%s AND user_id=%s",(filename,transaction_id,user["id"]))
    return jsonify({"message":"Receipt attached","filename":filename})


@enhancements.post("/api/import/message")
@token_required
def parse_message(user):
    text=(request.get_json(silent=True) or {}).get("text","")
    amount_match=re.search(r"(?:INR|Rs\.?|₹)\s*([\d,]+(?:\.\d{1,2})?)",text,re.I)
    date_match=re.search(r"(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})",text)
    typ="Income" if re.search(r"credited|received|deposit",text,re.I) else "Expense"
    merchant_match=(re.search(r"\bat\s+([A-Za-z0-9 .&_-]{2,50})",text,re.I)
                    or re.search(r"\bto\s+([A-Za-z0-9 .&_-]{2,50})",text,re.I)
                    or re.search(r"\bfrom\s+([A-Za-z0-9 .&_-]{2,50})",text,re.I))
    merchant=merchant_match.group(1).strip() if merchant_match else ""
    merchant=re.split(r"\s+on\s+\d",merchant,maxsplit=1,flags=re.I)[0].strip()
    return jsonify({"amount":float(amount_match.group(1).replace(",","")) if amount_match else None,
                    "transaction_date":date_match.group(1) if date_match else str(date.today()),
                    "type":typ,"merchant":merchant,
                    "description":text[:255],"requiresReview":True})


@enhancements.post("/api/accounts/<int:account_id>/reconcile")
@token_required
def reconcile(user,account_id):
    d=request.get_json(silent=True) or {}
    try: actual=float(d.get("actual_balance"))
    except (TypeError,ValueError): return jsonify({"error":"Enter the statement balance"}),400
    conn=get_conn(); cur=conn.cursor(dictionary=True)
    try:
        cur.execute("SELECT current_balance FROM accounts WHERE id=%s AND user_id=%s FOR UPDATE",(account_id,user["id"]))
        account=cur.fetchone()
        if not account: return jsonify({"error":"Account not found"}),404
        difference=round(actual-float(account["current_balance"]),2)
        cur.execute("UPDATE accounts SET current_balance=%s,last_reconciled_at=NOW() WHERE id=%s AND user_id=%s",(actual,account_id,user["id"]))
        if difference:
            cur.execute("""INSERT INTO transactions
                (user_id,account_id,type,category,amount,description,transaction_date,payment_method,review_status)
                VALUES(%s,%s,%s,'Adjustment',%s,'Account reconciliation adjustment',CURDATE(),'Reconciliation','Cleared')""",
                (user["id"],account_id,"Income" if difference>0 else "Expense",abs(difference)))
        conn.commit()
        return jsonify({"message":"Account reconciled","difference":difference,"adjustmentRecorded":bool(difference)})
    except Exception:
        conn.rollback(); return jsonify({"error":"Could not reconcile account"}),500
    finally:
        cur.close(); conn.close()


@enhancements.get("/api/connections")
@token_required
def connections(user):
    return jsonify(query("""SELECT c.*,a.account_name FROM account_connections c JOIN accounts a ON a.id=c.account_id
        WHERE c.user_id=%s ORDER BY c.id DESC""",(user["id"],),fetch=True))


@enhancements.post("/api/connections")
@token_required
def add_connection(user):
    d=request.get_json(silent=True) or {}
    query("""INSERT INTO account_connections(user_id,account_id,provider,status) VALUES(%s,%s,%s,'Setup Required')
        ON DUPLICATE KEY UPDATE provider=VALUES(provider)""",(user["id"],d.get("account_id"),d.get("provider","Open Banking Provider")))
    return jsonify({"message":"Connection created. Add provider credentials to activate automatic sync.",
                    "requiresProviderCredentials":True}),201


@enhancements.get("/api/notifications")
@token_required
def notifications(user):
    return jsonify(query("SELECT * FROM notifications WHERE user_id=%s ORDER BY created_at DESC LIMIT 50",(user["id"],),fetch=True))


@enhancements.patch("/api/notifications/<int:item_id>/read")
@token_required
def read_notification(user,item_id):
    query("UPDATE notifications SET is_read=1 WHERE id=%s AND user_id=%s",(item_id,user["id"]))
    return jsonify({"message":"Notification read"})


@enhancements.route("/api/preferences",methods=["GET","PUT"])
@token_required
def preferences(user):
    if request.method=="GET":
        row=query("SELECT * FROM user_preferences WHERE user_id=%s",(user["id"],),fetch=True,one=True)
        return jsonify(row or {"theme":"system","compact_mode":False,"reduced_motion":False,"dashboard_widgets":[]})
    d=request.get_json(silent=True) or {}; widgets=json.dumps(d.get("dashboard_widgets",[]))
    query("""INSERT INTO user_preferences(user_id,theme,compact_mode,reduced_motion,dashboard_widgets,report_email)
        VALUES(%s,%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE theme=VALUES(theme),
        compact_mode=VALUES(compact_mode),reduced_motion=VALUES(reduced_motion),
        dashboard_widgets=VALUES(dashboard_widgets),report_email=VALUES(report_email)""",
        (user["id"],d.get("theme","system"),bool(d.get("compact_mode")),bool(d.get("reduced_motion")),widgets,d.get("report_email")))
    return jsonify({"message":"Preferences saved"})


def report_rows(uid,start=None,end=None,category=None):
    where=["t.user_id=%s"]; params=[uid]
    if start: where.append("t.transaction_date>=%s"); params.append(start)
    if end: where.append("t.transaction_date<=%s"); params.append(end)
    if category: where.append("t.category=%s"); params.append(category)
    return query(f"""SELECT t.transaction_date,a.account_name,t.type,t.category,t.amount,t.description,
        t.payment_method,t.tags FROM transactions t JOIN accounts a ON a.id=t.account_id
        WHERE {' AND '.join(where)} ORDER BY t.transaction_date DESC""",tuple(params),fetch=True)


@enhancements.get("/api/reports/export/<fmt>")
@token_required
def export_report(user,fmt):
    if fmt not in ("csv","xlsx","pdf"): return jsonify({"error":"Choose csv, xlsx or pdf"}),400
    rows=report_rows(user["id"],request.args.get("start"),request.args.get("end"),request.args.get("category"))
    stamp=datetime.now().strftime("%Y%m%d")
    if fmt=="csv":
        out=io.StringIO(); fields=list(rows[0].keys()) if rows else ["message"]
        writer=csv.DictWriter(out,fieldnames=fields); writer.writeheader()
        if rows: writer.writerows(rows)
        data=io.BytesIO(out.getvalue().encode("utf-8-sig"))
    elif fmt=="xlsx":
        book=Workbook(); sheet=book.active; sheet.title="Transactions"
        fields=list(rows[0].keys()) if rows else ["message"]; sheet.append(fields)
        for row in rows: sheet.append([row[x] for x in fields])
        data=io.BytesIO(); book.save(data)
    else:
        data=io.BytesIO(); pdf=canvas.Canvas(data,pagesize=A4); width,height=A4
        pdf.setTitle("FinWise Financial Report"); y=height-45; pdf.setFont("Helvetica-Bold",16); pdf.drawString(40,y,"FinWise Financial Report"); y-=28
        pdf.setFont("Helvetica",8)
        for row in rows:
            line=f"{row['transaction_date']}  {row['type']:<8}  {row['category'] or 'Other':<14}  Rs {float(row['amount']):,.2f}  {(row['description'] or '')[:45]}"
            pdf.drawString(40,y,line); y-=14
            if y<40: pdf.showPage(); pdf.setFont("Helvetica",8); y=height-40
        pdf.save()
    data.seek(0)
    mime={"csv":"text/csv","xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","pdf":"application/pdf"}[fmt]
    return send_file(data,mimetype=mime,as_attachment=True,download_name=f"finwise_report_{stamp}.{fmt}")


@enhancements.route("/api/scheduled-reports",methods=["GET","POST"])
@token_required
def scheduled_reports(user):
    if request.method=="GET":
        return jsonify(query("SELECT * FROM scheduled_reports WHERE user_id=%s ORDER BY id DESC",(user["id"],),fetch=True))
    d=request.get_json(silent=True) or {}; frequency=d.get("frequency","Monthly")
    next_send=date.today()+timedelta(days=7 if frequency=="Weekly" else 30)
    query("INSERT INTO scheduled_reports(user_id,frequency,format,email,next_send) VALUES(%s,%s,%s,%s,%s)",
          (user["id"],frequency,d.get("format","pdf"),d.get("email"),next_send))
    return jsonify({"message":"Report reminder saved locally. Automatic email delivery is unavailable until an SMTP provider is connected.","requiresEmailProvider":True}),201


def membership(uid):
    return query("""SELECT hm.*,h.name,h.owner_user_id FROM household_members hm
        JOIN households h ON h.id=hm.household_id WHERE hm.user_id=%s AND hm.status='Active' LIMIT 1""",(uid,),fetch=True,one=True)


@enhancements.route("/api/household",methods=["GET","POST"])
@token_required
def household(user):
    if request.method=="GET":
        member=membership(user["id"])
        if not member: return jsonify(None)
        members=query("""SELECT hm.id,hm.role,hm.status,u.name,u.email FROM household_members hm
            JOIN users u ON u.id=hm.user_id WHERE hm.household_id=%s""",(member["household_id"],),fetch=True)
        invites=query("SELECT id,email,role,status,created_at,expires_at,accepted_at,revoked_at FROM household_invites WHERE household_id=%s ORDER BY id DESC",(member["household_id"],),fetch=True)
        return jsonify({"household":member,"members":members,"invites":invites})
    if membership(user["id"]): return jsonify({"error":"You already belong to a household"}),409
    name=(request.get_json(silent=True) or {}).get("name","My Household").strip()[:120]
    conn=get_conn(); cur=conn.cursor()
    try:
        cur.execute("INSERT INTO households(name,owner_user_id) VALUES(%s,%s)",(name,user["id"])); hid=cur.lastrowid
        cur.execute("INSERT INTO household_members(household_id,user_id,role) VALUES(%s,%s,'Owner')",(hid,user["id"]))
        conn.commit()
    finally: cur.close(); conn.close()
    return jsonify({"message":"Household created"}),201


@enhancements.post("/api/household/invites")
@token_required
def invite_household(user):
    member=membership(user["id"]); d=request.get_json(silent=True) or {}
    if not member or member["role"] not in ("Owner","Admin"): return jsonify({"error":"Admin access required"}),403
    email=(d.get("email") or "").strip().lower()
    role=d.get("role","Member")
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+",email): return jsonify({"error":"Enter a valid member email"}),400
    if role not in ("Admin","Member","Viewer"): return jsonify({"error":"Choose a valid role"}),400
    token=str(uuid.uuid4())
    expires=datetime.now()+timedelta(days=7)
    query("INSERT INTO household_invites(household_id,email,role,token,expires_at) VALUES(%s,%s,%s,%s,%s)",
          (member["household_id"],email,role,token,expires))
    return jsonify({"message":"Invite created. Share this token once; it expires in 7 days.","inviteToken":token,"expiresAt":expires}),201


@enhancements.get("/api/household/invites/mine")
@token_required
def my_household_invites(user):
    if membership(user["id"]): return jsonify([])
    email=str(user.get("email") or "").strip().lower()
    invites=query("""SELECT hi.id,hi.email,hi.role,hi.status,hi.created_at,hi.expires_at,h.name household_name
        FROM household_invites hi JOIN households h ON h.id=hi.household_id
        WHERE LOWER(hi.email)=LOWER(%s) AND hi.status='Pending'
        AND (hi.expires_at IS NULL OR hi.expires_at>NOW())
        ORDER BY hi.id DESC""",(email,),fetch=True)
    return jsonify(invites)


def join_household_from_invite(user, invite):
    if invite.get("email") and str(invite["email"]).lower()!=str(user["email"]).lower():
        return jsonify({"error":"This invite was created for a different email"}),403
    conn=get_conn(); cur=conn.cursor()
    try:
        cur.execute("INSERT INTO household_members(household_id,user_id,role,status) VALUES(%s,%s,%s,'Active')",
                    (invite["household_id"],user["id"],invite["role"]))
        cur.execute("UPDATE household_invites SET status='Accepted',accepted_at=NOW() WHERE id=%s",(invite["id"],))
        conn.commit()
    except Exception:
        conn.rollback(); return jsonify({"error":"Could not join family group"}),500
    finally:
        cur.close(); conn.close()
    return jsonify({"message":"Joined family group"})


@enhancements.post("/api/household/invites/accept")
@token_required
def accept_household_invite(user):
    if membership(user["id"]): return jsonify({"error":"You already belong to a family group"}),409
    token=(request.get_json(silent=True) or {}).get("token","").strip()
    if not token: return jsonify({"error":"Enter invite token"}),400
    invite=query("SELECT * FROM household_invites WHERE token=%s AND status='Pending' AND (expires_at IS NULL OR expires_at>NOW())",(token,),fetch=True,one=True)
    if not invite: return jsonify({"error":"Invite token is invalid or already used"}),404
    return join_household_from_invite(user,invite)


@enhancements.post("/api/household/invites/<int:invite_id>/accept")
@token_required
def accept_household_invite_by_id(user,invite_id):
    if membership(user["id"]): return jsonify({"error":"You already belong to a family group"}),409
    invite=query("""SELECT * FROM household_invites
        WHERE id=%s AND status='Pending' AND (expires_at IS NULL OR expires_at>NOW())""",(invite_id,),fetch=True,one=True)
    if not invite: return jsonify({"error":"Invite is invalid, expired or already used"}),404
    return join_household_from_invite(user,invite)


@enhancements.patch("/api/household/invites/<int:invite_id>/revoke")
@token_required
def revoke_household_invite(user,invite_id):
    member=membership(user["id"])
    if not member or member["role"] not in ("Owner","Admin"): return jsonify({"error":"Admin access required"}),403
    query("UPDATE household_invites SET status='Revoked',revoked_at=NOW() WHERE id=%s AND household_id=%s AND status='Pending'",
          (invite_id,member["household_id"]))
    return jsonify({"message":"Invite revoked"})


@enhancements.route("/api/approvals",methods=["GET","POST"])
@token_required
def approvals(user):
    member=membership(user["id"])
    if not member: return jsonify([]) if request.method=="GET" else (jsonify({"error":"Join a household first"}),400)
    if request.method=="GET":
        return jsonify(query("""SELECT p.*,u.name requested_by_name FROM transaction_approvals p
            JOIN users u ON u.id=p.requested_by WHERE p.household_id=%s ORDER BY p.id DESC""",(member["household_id"],),fetch=True))
    d=request.get_json(silent=True) or {}
    query("INSERT INTO transaction_approvals(household_id,transaction_id,requested_by,note) VALUES(%s,%s,%s,%s)",
          (member["household_id"],d.get("transaction_id"),user["id"],d.get("note")))
    return jsonify({"message":"Approval requested"}),201


@enhancements.patch("/api/approvals/<int:item_id>")
@token_required
def review_approval(user,item_id):
    member=membership(user["id"]); status=(request.get_json(silent=True) or {}).get("status")
    if not member or member["role"] not in ("Owner","Admin") or status not in ("Approved","Rejected"):
        return jsonify({"error":"Admin access and a valid decision are required"}),403
    query("UPDATE transaction_approvals SET status=%s,reviewed_by=%s WHERE id=%s AND household_id=%s",
          (status,user["id"],item_id,member["household_id"]))
    return jsonify({"message":f"Request {status.lower()}"})
